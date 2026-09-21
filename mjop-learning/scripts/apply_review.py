#!/usr/bin/env python3
"""
apply_review.py  (stap: ingevuld review-Excel -> data/verified/)

Leest het door een mens ingevulde reports/review_batch1.xlsx (zie
export_review_sheet.py) terug en schrijft de resulterende, door een mens
gecontroleerde records naar data/verified/*.json.

Per rij met een ingevulde BESLISSING (accept/edit/reject, zie CLAUDE.md):
  - accept: de huidige waarde (vaak normalized_value=null, of een afgeleide
    unit_cost_calculated) is correct zoals hij is - de post blijft zo, maar
    is nu een bevestigd, geen onopgelost, oordeel. requires_human_review
    wordt False.
  - edit:   GECORRIGEERDE_WAARDE vervangt normalized_value van het veld dat
    de "veld"-kolom aanwijst (actie of eenheid). Alleen zinvol als "veld" =
    actie/eenheid; bij "veld" = geen (kostenconflict/afgeleide prijs, geen
    vocabulaire-veld om te corrigeren) wordt edit geweigerd. Vereist een
    ingevulde GECORRIGEERDE_WAARDE, anders wordt de rij overgeslagen (nooit
    zelf een waarde verzinnen).
  - reject: de post blijft ongewijzigd en requires_human_review blijft
    True - dit betekent "een mens heeft ernaar gekeken en ook geen
    betrouwbaar antwoord", niet "opgelost". NOTITIES legt vast waarom.

Voor "veld" = actie/eenheid/defect wordt human_verification vastgelegd op
het betreffende vocabulaire-paar (action["action"]/action["unit"]/
obs["defect"]), net als het patroon in schemas/_human_verification.schema.json.
Voor ELKE beoordeelde rij (ook "veld" = geen, bijv. bij een kostenconflict)
wordt bovendien altijd een leesbare review_note op de action/observation
zelf gezet, zodat er ook een auditspoor is voor beslissingen zonder
vocabulaire-veld.

Rijen zonder ingevulde BESLISSING worden overgeslagen (nog niet
beoordeeld). Documenten zonder enige beoordeelde rij krijgen geen
data/verified/-bestand - "verified" betekent hier "voor zover beoordeeld",
niet "elk veld is met de hand gecontroleerd" (zie CLAUDE.md: geen claims
zonder meting - evaluate_dataset.py's cijfer is dus alleen zinvol voor de
velden die daadwerkelijk beoordeeld zijn).

Gebruik:
    python3 scripts/apply_review.py --reviewer "jouw@email" \
        --xlsx reports/review_batch1.xlsx
"""
import argparse
import glob
import json
import os
from datetime import datetime, timezone

from openpyxl import load_workbook

VALID_DECISIONS = {"accept", "edit", "reject"}


def now_iso():
    return datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


def apply_pair_decision(pair, decision, corrected_value, reviewer, timestamp, notes=None):
    """Past een ACCEPT/EDIT/REJECT-beslissing toe op een {original_value,
    normalized_value}-pair en legt human_verification vast. Retourneert
    (pair, warning_or_None). Verzint nooit een waarde: bij edit zonder
    corrected_value gebeurt niets en komt er een waarschuwing terug."""
    decision = (decision or "").strip().lower()
    if decision not in VALID_DECISIONS:
        return pair, f"onbekende beslissing {decision!r}, rij overgeslagen"

    verification = {
        "status": decision,
        "original_ai_value": pair.get("normalized_value"),
        "reviewer": reviewer,
        "timestamp": timestamp,
    }
    if notes:
        verification["notes"] = str(notes)

    if decision == "edit":
        if not corrected_value:
            return pair, "edit zonder GECORRIGEERDE_WAARDE, rij overgeslagen (geen aanname ingevuld)"
        verification["edited_value"] = corrected_value
        pair["normalized_value"] = corrected_value
        pair["requires_human_review"] = False
    elif decision == "accept":
        pair["requires_human_review"] = False
    # reject: normalized_value en requires_human_review blijven ongewijzigd

    pair["human_verification"] = verification
    return pair, None


def apply_review_note(container, decision, reviewer, timestamp, notes=None):
    """Zet een leesbaar review_note op de maintenance_action/observation
    zelf - ook nuttig als er geen vocabulaire-paar bij deze rij hoort
    (bijv. een kostenconflict of een afgeleide unit_cost_calculated)."""
    note = f"[{decision} door {reviewer} op {timestamp}]"
    if notes:
        note += f" {notes}"
    container["review_note"] = note


def read_rows(ws, columns):
    header = [c.value for c in ws[1]]
    idx = {name: header.index(name) for name in columns if name in header}
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[idx["BESLISSING"]] is None or str(row[idx["BESLISSING"]]).strip() == "":
            continue
        rows.append({name: row[i] for name, i in idx.items()})
    return rows


def process_action_row(action, row, reviewer, timestamp):
    decision = str(row["BESLISSING"]).strip().lower()
    field = (row.get("veld") or "geen").strip().lower()
    warning = None
    new_value = None

    if field in ("actie", "eenheid"):
        pair_key = "action" if field == "actie" else "unit"
        pair = action.get(pair_key) or {}
        pair, warning = apply_pair_decision(pair, decision, row.get("GECORRIGEERDE_WAARDE"), reviewer, timestamp, row.get("NOTITIES"))
        action[pair_key] = pair
        if not warning and decision == "edit":
            new_value = pair.get("normalized_value")
    elif decision == "edit":
        warning = "edit zonder vocabulaire-veld (veld=geen) - GECORRIGEERDE_WAARDE kan niet toegepast worden, gebruik accept/reject"

    if not warning:
        apply_review_note(action, decision, reviewer, timestamp, row.get("NOTITIES"))
        if decision in ("accept", "edit") and not (action.get("action") or {}).get("requires_human_review") and not (action.get("unit") or {}).get("requires_human_review"):
            action["requires_human_review"] = False
    return warning, new_value


def process_observation_row(obs, row, reviewer, timestamp):
    decision = str(row["BESLISSING"]).strip().lower()
    defect = obs.get("defect") or {}
    defect, warning = apply_pair_decision(defect, decision, row.get("GECORRIGEERDE_WAARDE"), reviewer, timestamp, row.get("NOTITIES"))
    obs["defect"] = defect
    new_value = defect.get("normalized_value") if not warning and decision == "edit" else None

    if not warning:
        apply_review_note(obs, decision, reviewer, timestamp, row.get("NOTITIES"))
        if decision in ("accept", "edit") and not defect.get("requires_human_review"):
            obs["requires_human_review"] = False
    return warning, new_value


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", default="reports/review_batch1.xlsx")
    ap.add_argument("--normalized-dir", default="data/normalized")
    ap.add_argument("--verified-dir", default="data/verified")
    ap.add_argument("--reviewer", required=True)
    args = ap.parse_args()

    wb = load_workbook(args.xlsx, data_only=True)
    action_rows = read_rows(wb["maintenance_actions"], [
        "document_id", "action_id", "veld", "BESLISSING", "GECORRIGEERDE_WAARDE", "NOTITIES",
    ])
    obs_rows = read_rows(wb["observations"], [
        "document_id", "observation_id", "BESLISSING", "GECORRIGEERDE_WAARDE", "NOTITIES",
    ])

    timestamp = now_iso()
    records = {}
    warnings = []
    new_normalized_values = set()

    def get_record(doc_id):
        if doc_id not in records:
            path = os.path.join(args.normalized_dir, f"{doc_id}.json")
            records[doc_id] = json.load(open(path))
        return records[doc_id]

    for row in action_rows:
        rec = get_record(row["document_id"])
        action = next((a for a in rec.get("maintenance_actions", []) if a.get("action_id") == row["action_id"]), None)
        if action is None:
            warnings.append(f"{row['document_id']}/{row['action_id']}: action_id niet gevonden, overgeslagen")
            continue
        warning, new_value = process_action_row(action, row, args.reviewer, timestamp)
        if warning:
            warnings.append(f"{row['document_id']}/{row['action_id']}: {warning}")
        elif new_value:
            new_normalized_values.add(new_value)

    for row in obs_rows:
        rec = get_record(row["document_id"])
        obs = next((o for o in rec.get("observations", []) if o.get("observation_id") == row["observation_id"]), None)
        if obs is None:
            warnings.append(f"{row['document_id']}/{row['observation_id']}: observation_id niet gevonden, overgeslagen")
            continue
        warning, new_value = process_observation_row(obs, row, args.reviewer, timestamp)
        if warning:
            warnings.append(f"{row['document_id']}/{row['observation_id']}: {warning}")
        elif new_value:
            new_normalized_values.add(new_value)

    os.makedirs(args.verified_dir, exist_ok=True)
    for doc_id, rec in records.items():
        out_path = os.path.join(args.verified_dir, f"{doc_id}.json")
        json.dump(rec, open(out_path, "w"), ensure_ascii=False, indent=2)

    print(f"{len(records)} document(en) geschreven naar {args.verified_dir}/")
    print(f"{len(action_rows)} maintenance_action-rijen en {len(obs_rows)} observation-rijen verwerkt.")
    if new_normalized_values:
        print(
            "Nieuwe normalized_values via EDIT geintroduceerd (nog NIET automatisch aan vocabularies/ toegevoegd - "
            "controleer en voeg zelf toe als dit een herbruikbare term is): " + ", ".join(sorted(new_normalized_values))
        )
    if warnings:
        print(f"\n{len(warnings)} waarschuwing(en):")
        for w in warnings:
            print(f"  - {w}")


if __name__ == "__main__":
    main()
