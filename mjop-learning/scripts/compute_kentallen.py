#!/usr/bin/env python3
"""
compute_kentallen.py  (stap: VERIFIED -> KENTALLEN)

Berekent kentallen (eenheidsprijs-benchmarks) uit data/verified/*.json,
gegroepeerd per (element_code, actie, eenheid) - dus op de manier waarop
wij onze eigen MJOP's opbouwen (vocabularies/element_code.json, ons eigen
NL-SfB-achtige coderingssysteem), niet een generieke vrij-tekst-indeling.
Documenten zonder element_code (bijv. DOC-004, van Innax, niet van
onszelf) vallen hierdoor vanzelf buiten deze berekening.

Alleen posten die de pipeline betrouwbaar genoeg acht tellen mee:
  - requires_human_review moet False zijn op de actie EN op het gekoppelde
    element (nooit gevlagd, of na een menselijke ACCEPT/EDIT weer
    opgelost - een REJECT of nog-onbeoordeelde post telt niet mee).
  - cost_conflict moet False zijn (bij een conflict weten we niet welk
    bedrag klopt - geen kental uit een onopgeloste tegenstrijdigheid).
  - element_code.normalized_value, action.normalized_value en
    unit.normalized_value moeten alle drie gevalideerd zijn (zonder die
    combinatie is een prijs niet herbruikbaar/vergelijkbaar).

Voor elke groep wordt APART geteld/gerapporteerd wat uit een letterlijke
documentprijs komt (unit_cost.value) en wat uit een afgeleide prijs
(unit_cost_calculated = total_cost_as_stated / hoeveelheid) - nooit
stilzwijgend gemengd, want een afgeleide prijs is per definitie minder
betrouwbaar (zie schemas/maintenance_action.schema.json).

Prijzen zijn NIET geindexeerd naar een gemeenschappelijk prijspeil - elke
groep vermeldt welke cost_year(en)/prijspeiljaren zijn gezien, zodat nooit
een oude prijs stilzwijgend als actuele marktprijs wordt gepresenteerd
(zie CLAUDE.md).

Gebruik:
    python3 scripts/compute_kentallen.py
"""
import argparse
import glob
import json
import os
import statistics
import sys
from decimal import Decimal

sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "scripts"))
from normalize_batch import to_decimal  # noqa: E402

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def load_element_code_meta(vocab_dir):
    doc = json.load(open(os.path.join(vocab_dir, "element_code.json")))
    hoofdgroepen = doc.get("hoofdgroepen", {})
    meta = {}
    for entry in doc.get("entries", []):
        code = entry["normalized_value"]
        hg_code = entry.get("hoofdgroep_code")
        meta[code] = {
            "label_nl": entry.get("label_nl"),
            "hoofdgroep_code": hg_code,
            "hoofdgroep_label": hoofdgroepen.get(hg_code),
        }
    return meta


def is_reliable(action, element):
    if action.get("requires_human_review"):
        return False
    if action.get("cost_conflict"):
        return False
    if (element.get("element_code") or {}).get("requires_human_review"):
        return False
    return True


def collect_price_points(verified_dir):
    """Loopt alle data/verified/*.json door en geeft per bruikbare actie een
    dict terug met de sleutel (element_code, action, unit) en het type/de
    waarde van de prijs. Telt ook bij waarom posten worden overgeslagen."""
    points = []
    skipped = {
        "requires_human_review_of_conflict": 0,
        "geen_element_code": 0,
        "geen_actie_of_eenheid": 0,
        "geen_prijs": 0,
    }

    for path in sorted(glob.glob(os.path.join(verified_dir, "*.json"))):
        rec = json.load(open(path))
        element_by_id = {el.get("element_id"): el for el in rec.get("elements", [])}

        for action in rec.get("maintenance_actions", []):
            element = element_by_id.get(action.get("element_id"))
            if element is None:
                continue
            if not is_reliable(action, element):
                skipped["requires_human_review_of_conflict"] += 1
                continue

            element_code = (element.get("element_code") or {}).get("normalized_value")
            if not element_code:
                skipped["geen_element_code"] += 1
                continue

            action_value = (action.get("action") or {}).get("normalized_value")
            unit_value = (action.get("unit") or {}).get("normalized_value")
            if not action_value or not unit_value:
                skipped["geen_actie_of_eenheid"] += 1
                continue

            unit_cost = to_decimal((action.get("unit_cost") or {}).get("value"))
            unit_cost_calc = to_decimal(action.get("unit_cost_calculated"))
            if unit_cost is not None:
                price, price_type = unit_cost, "literal"
            elif unit_cost_calc is not None:
                price, price_type = unit_cost_calc, "calculated"
            else:
                skipped["geen_prijs"] += 1
                continue

            points.append({
                "element_code": element_code,
                "action": action_value,
                "unit": unit_value,
                "price": price,
                "price_type": price_type,
                "cost_year": action.get("cost_year"),
                "document_id": rec.get("document_id"),
                "action_id": action.get("action_id"),
            })

    return points, skipped


def stats_for(prices):
    if not prices:
        return None
    prices = sorted(prices)
    return {
        "n": len(prices),
        "min": str(min(prices).quantize(Decimal("0.01"))),
        "max": str(max(prices).quantize(Decimal("0.01"))),
        "mean": str((sum(prices) / len(prices)).quantize(Decimal("0.01"))),
        "median": str(Decimal(str(statistics.median(prices))).quantize(Decimal("0.01"))),
    }


def aggregate(points, element_code_meta):
    groups = {}
    for p in points:
        key = (p["element_code"], p["action"], p["unit"])
        groups.setdefault(key, {"literal": [], "calculated": [], "cost_years": set(), "documents": set(), "action_ids": []})
        g = groups[key]
        g["literal" if p["price_type"] == "literal" else "calculated"].append(p["price"])
        if p["cost_year"]:
            g["cost_years"].add(p["cost_year"])
        g["documents"].add(p["document_id"])
        g["action_ids"].append(p["action_id"])

    results = []
    for (element_code, action, unit), g in groups.items():
        meta = element_code_meta.get(element_code, {})
        results.append({
            "element_code": element_code,
            "element_code_label": meta.get("label_nl"),
            "hoofdgroep_code": meta.get("hoofdgroep_code"),
            "hoofdgroep_label": meta.get("hoofdgroep_label"),
            "action": action,
            "unit": unit,
            "literal": stats_for(g["literal"]),
            "calculated": stats_for(g["calculated"]),
            "cost_years": sorted(g["cost_years"]),
            "n_documents": len(g["documents"]),
            "documents": sorted(g["documents"]),
        })

    results.sort(key=lambda r: (r["hoofdgroep_code"] or "", r["element_code"], r["action"], r["unit"]))
    return results


def write_xlsx(results, out_path):
    columns = [
        "hoofdgroep_code", "hoofdgroep_label", "element_code", "element_code_label",
        "action", "unit",
        "n_literal", "literal_min", "literal_median", "literal_max",
        "n_calculated", "calculated_min", "calculated_median", "calculated_max",
        "cost_years", "n_documents",
    ]
    wb = Workbook()
    ws = wb.active
    ws.title = "kentallen"
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in results:
        lit = r["literal"] or {}
        calc = r["calculated"] or {}
        ws.append([
            r["hoofdgroep_code"], r["hoofdgroep_label"], r["element_code"], r["element_code_label"],
            r["action"], r["unit"],
            lit.get("n", 0), lit.get("min"), lit.get("median"), lit.get("max"),
            calc.get("n", 0), calc.get("min"), calc.get("median"), calc.get("max"),
            ", ".join(str(y) for y in r["cost_years"]), r["n_documents"],
        ])
    for i, col in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(40, len(col) + 2))
    ws.freeze_panes = "A2"
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--verified-dir", default="data/verified")
    ap.add_argument("--vocab-dir", default="vocabularies")
    ap.add_argument("--out-json", default="data/kentallen/kentallen_batch1.json")
    ap.add_argument("--out-xlsx", default="reports/kentallen_batch1.xlsx")
    args = ap.parse_args()

    element_code_meta = load_element_code_meta(args.vocab_dir)
    points, skipped = collect_price_points(args.verified_dir)
    results = aggregate(points, element_code_meta)

    os.makedirs(os.path.dirname(args.out_json), exist_ok=True)
    json.dump(
        {
            "note": (
                "Kentallen per (element_code, actie, eenheid), NIET geindexeerd naar een "
                "gemeenschappelijk prijspeil - zie cost_years per groep. 'literal' = letterlijke "
                "documentprijs (unit_cost.value), 'calculated' = afgeleid uit total/hoeveelheid "
                "(minder betrouwbaar, zie schemas/maintenance_action.schema.json). Gebaseerd op "
                "data/verified/ (na menselijke review); documenten zonder element_code (niet van "
                "onszelf, bijv. DOC-004/Innax) tellen niet mee."
            ),
            "groups": results,
        },
        open(args.out_json, "w"),
        ensure_ascii=False,
        indent=2,
    )
    write_xlsx(results, args.out_xlsx)

    n_points = len(points)
    n_groups = len(results)
    print(f"{n_points} prijspunten verwerkt in {n_groups} kentallen-groepen (element_code x actie x eenheid).")
    print(f"-> {args.out_json}")
    print(f"-> {args.out_xlsx}")
    print()
    print("Overgeslagen posten (nooit meegeteld in een kental):")
    for reason, n in skipped.items():
        print(f"  {reason}: {n}")


if __name__ == "__main__":
    main()
