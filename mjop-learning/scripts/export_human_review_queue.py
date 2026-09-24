#!/usr/bin/env python3
"""
export_human_review_queue.py  (COMPARABILITY -> REVIEWQUEUE VOOR EEN MENS)

Bouwt reproduceerbaar de eerste reviewqueue (docs/human_review_v1.md) uit
data/comparability/comparability_batch1.json en de genormaliseerde price
observations, en schrijft reports/human_review_queue_v1.xlsx.

Selectie (queue v1):
  - systeemklasse COMPARABLE_WITH_CAVEATS;
  - beide observations independent_input = true (sluit o.a. DOC-005
    "(uitgevoerd JJJJ)" en POSSIBLY_DEPENDENT uit);
  - geen DOC-001-observation zonder materiaal (nog onopgeloste
    DOC-001-materiaalgevallen);
  - geen UNKNOWN-paren.

De Excel is een reviewinstrument, niet de source of truth: de kolommen voor
de menselijke beslissing zijn leeg en dit script maakt nooit een beslissing
aan. Leest alleen; bron-, normalisatie- en comparability-data blijven
ongewijzigd. Geen scores, confidence of rangschikking.

Bevat ook de controles op de invarianten van de beslissingenopslag die het
JSON-schema niet afdwingt (store_invariant_errors, append_only_errors).

Gebruik:
    python scripts/export_human_review_queue.py [--dry-run]
"""
import argparse
import hashlib
import io
import json
import os
import re
import zipfile
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

QUEUE_VERSION = "human_review_queue_v1"
DECISION_OPTIONS = ["COMPARABLE", "COMPARABLE_WITH_CAVEATS", "NOT_COMPARABLE", "UNKNOWN"]
SELECTION = [
    "systeemklasse COMPARABLE_WITH_CAVEATS",
    "beide observations independent_input = true",
    "geen DOC-001-observation zonder materiaal (material.source leeg)",
    "geen UNKNOWN-paren",
]
SIDE_COLUMNS = [
    "observation_id", "document_id", "source_cluster", "page", "line", "source_text",
    "element_code_original", "object_description", "action_text", "material", "material_source",
    "quantity", "unit_original", "derived_price_per_execution", "executions_in_window",
    "price_level_date", "vat_basis",
]
PAIR_COLUMNS = ["pair_id", "element_code", "action", "unit", "system_class", "pair_caveats",
                "hard_violations", "unknown_reasons", "observation_caveats_a", "observation_caveats_b"]
HUMAN_COLUMNS = ["HUMAN_DECISION", "DECISION_REASON", "DECISION_CAVEATS", "REVIEWER", "NOTES"]
FIXED_TIMESTAMP = datetime(2000, 1, 1)


def sha256_file(path):
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for chunk in iter(lambda: f.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def store_invariant_errors(store):
    """Invarianten over records heen die het JSON-schema niet afdwingt
    (docs/human_review_v1.md): unieke decision_id, per pair_id hoogstens één
    ACTIVE record, en 'supersedes' verwijst naar een bestaand record van
    hetzelfde paar dat SUPERSEDED is."""
    errors, by_id = [], {}
    for r in store["records"]:
        if r["decision_id"] in by_id:
            errors.append(f"{r['decision_id']}: decision_id komt meer dan eens voor")
        by_id[r["decision_id"]] = r
    active = {}
    for r in store["records"]:
        if r["status"] == "ACTIVE":
            active.setdefault(r["pair_id"], []).append(r["decision_id"])
        old = by_id.get(r["supersedes"]) if r["supersedes"] else None
        if r["supersedes"] and old is None:
            errors.append(f"{r['decision_id']}: supersedes verwijst naar onbekende {r['supersedes']}")
        elif old is not None and (old["pair_id"] != r["pair_id"] or old["status"] != "SUPERSEDED"):
            errors.append(f"{r['decision_id']}: vervangen record {old['decision_id']} moet hetzelfde paar "
                          f"hebben en SUPERSEDED zijn")
    errors += [f"{pid}: meer dan één ACTIVE record ({', '.join(ids)})" for pid, ids in sorted(active.items())
               if len(ids) > 1]
    return errors


def append_only_errors(old_store, new_store):
    """Een nieuwe versie van de opslag mag alleen records toevoegen; een bestaand
    record blijft ongewijzigd, behalve status ACTIVE -> SUPERSEDED."""
    new = {r["decision_id"]: r for r in new_store["records"]}
    errors = []
    for r in old_store["records"]:
        n = new.get(r["decision_id"])
        if n is None:
            errors.append(f"{r['decision_id']}: record verwijderd")
            continue
        if {k: v for k, v in n.items() if k != "status"} != {k: v for k, v in r.items() if k != "status"}:
            errors.append(f"{r['decision_id']}: record overschreven")
        if n["status"] != r["status"] and (r["status"], n["status"]) != ("ACTIVE", "SUPERSEDED"):
            errors.append(f"{r['decision_id']}: statuswijziging {r['status']} -> {n['status']} niet toegestaan")
    return errors


def unresolved_doc001_material(assessment):
    return assessment["document_id"] == "DOC-001" and not (assessment.get("material") or {}).get("source")


def select_pairs(comparability):
    """Paren voor queue v1, gesorteerd op pair_id."""
    obs = {a["observation_id"]: a for a in comparability["observations"]}
    out = []
    for p in comparability["pairs"]:
        if p["class"] != "COMPARABLE_WITH_CAVEATS":
            continue
        sides = [obs[i] for i in p["observation_ids"]]
        if not all(a["independent_input"] for a in sides):
            continue
        if any(unresolved_doc001_material(a) for a in sides):
            continue
        out.append(p)
    return sorted(out, key=lambda p: p["pair_id"])


def side_values(assessment, normalized):
    primary = next(r for r in normalized["source_ref"]["source_representations"]
                   if r["role"] == "primary_financial_row")
    derived = assessment["derived_unit_price_per_execution"] or {}
    material = assessment.get("material") or {}
    return {
        "observation_id": assessment["observation_id"],
        "document_id": assessment["document_id"],
        "source_cluster": assessment["source_cluster"],
        "page": primary["page"],
        "line": primary["line"],
        "source_text": primary["source_text"],
        "element_code_original": normalized["element"]["element_code_original"],
        "object_description": normalized["element"]["element_description_original"],
        "action_text": normalized["action"]["action_text_original"],
        "material": material.get("original"),
        "material_source": material.get("source"),
        "quantity": normalized["price"]["quantity_value"],
        "unit_original": normalized["unit"]["unit_original"],
        "derived_price_per_execution": derived.get("value"),
        "executions_in_window": derived.get("executions_in_window"),
        "price_level_date": normalized["price"]["price_level_date"],
        "vat_basis": normalized["price"]["vat_basis"],
    }


def build_rows(comparability, normalized_doc):
    obs = {a["observation_id"]: a for a in comparability["observations"]}
    norm = {o["observation_id"]: o for o in normalized_doc["observations"]}
    rows = []
    for p in select_pairs(comparability):
        a_id, b_id = p["observation_ids"]
        row = {
            "pair_id": p["pair_id"],
            "element_code": p["candidate_key"][0],
            "action": p["candidate_key"][1],
            "unit": p["candidate_key"][2],
            "system_class": p["class"],
            "pair_caveats": ", ".join(p["pair_caveats"]),
            "hard_violations": ", ".join(p["hard_violations"]),
            "unknown_reasons": ", ".join(p["unknown_reasons"]),
            "observation_caveats_a": ", ".join(p["observation_caveats"]["a"]),
            "observation_caveats_b": ", ".join(p["observation_caveats"]["b"]),
        }
        for prefix, oid in (("a", a_id), ("b", b_id)):
            for k, v in side_values(obs[oid], norm[oid]).items():
                row[f"{prefix}_{k}"] = v
        for k in HUMAN_COLUMNS:
            row[k] = None  # bewust leeg: alleen een mens vult dit in
        rows.append(row)
    return rows


def columns():
    return PAIR_COLUMNS + [f"a_{c}" for c in SIDE_COLUMNS] + [f"b_{c}" for c in SIDE_COLUMNS] + HUMAN_COLUMNS


def write_xlsx(rows, meta, path):
    wb = Workbook()
    ws = wb.active
    ws.title = "queue"
    cols = columns()
    ws.append(cols)
    for c in ws[1]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([r[c] for c in cols])
    for i, c in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = 14 if c in HUMAN_COLUMNS else 18
    ws.freeze_panes = "B2"
    if rows:
        col = get_column_letter(cols.index("HUMAN_DECISION") + 1)
        dv = DataValidation(type="list", formula1='"' + ",".join(DECISION_OPTIONS) + '"', allow_blank=True)
        dv.add(f"{col}2:{col}{len(rows) + 1}")
        ws.add_data_validation(dv)

    info = wb.create_sheet("toelichting")
    lines = [
        ("queue_version", meta["queue_version"]),
        ("rule_version", meta["rule_version"]),
        ("pairs_in_queue", meta["pairs_in_queue"]),
        ("comparability_output_sha256", meta["comparability_output_sha256"]),
        ("normalized_observations_sha256", meta["normalized_observations_sha256"]),
        ("status", "Reviewinstrument, NIET de source of truth. Beslissingen horen als records in "
                   "data/review_decisions/human_decision_records.json (docs/human_review_v1.md)."),
        ("HUMAN_DECISION", " / ".join(DECISION_OPTIONS)),
        ("DECISION_REASON", "verplicht bij elke beslissing"),
    ] + [("selectie", s) for s in meta["selection"]]
    for k, v in lines:
        info.append([k, v])
    info.column_dimensions["A"].width = 32
    info.column_dimensions["B"].width = 100

    wb.properties.created = FIXED_TIMESTAMP
    buf = io.BytesIO()
    wb.save(buf)
    # tijdstempels vastzetten (openpyxl zet 'modified' bij opslaan op nu, zip-entries
    # krijgen de huidige tijd) zodat dezelfde invoer byte-identieke output geeft
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                              rb"\g<1>" + FIXED_TIMESTAMP.strftime("%Y-%m-%dT%H:%M:%SZ").encode() + rb"\g<2>", data)
            info_ = zipfile.ZipInfo(item.filename, date_time=(1980, 1, 1, 0, 0, 0))
            info_.compress_type = zipfile.ZIP_DEFLATED
            info_.external_attr = item.external_attr
            dst.writestr(info_, data)
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "wb") as f:
        f.write(out.getvalue())


def build(project_root):
    cmp_path = os.path.join(project_root, "data", "comparability", "comparability_batch1.json")
    norm_path = os.path.join(project_root, "data", "price_observations", "price_observations_batch1_normalized.json")
    comparability = json.load(open(cmp_path, encoding="utf-8"))
    normalized = json.load(open(norm_path, encoding="utf-8"))
    rows = build_rows(comparability, normalized)
    meta = {
        "queue_version": QUEUE_VERSION,
        "rule_version": comparability["rules_version"],
        "pairs_in_queue": len(rows),
        "comparability_output_sha256": sha256_file(cmp_path),
        "normalized_observations_sha256": sha256_file(norm_path),
        "selection": SELECTION,
    }
    return rows, meta


def main():
    root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=os.path.join("reports", "human_review_queue_v1.xlsx"))
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()
    rows, meta = build(root)
    print(json.dumps({k: v for k, v in meta.items() if k != "selection"}, indent=2))
    for r in rows:
        print(" ", r["pair_id"], r["element_code"], r["action"], r["unit"], "|", r["pair_caveats"])
    if args.dry_run:
        print("\n--dry-run: niets geschreven.")
        return
    write_xlsx(rows, meta, os.path.join(root, args.out))
    print(f"\n-> {args.out}")


if __name__ == "__main__":
    main()
