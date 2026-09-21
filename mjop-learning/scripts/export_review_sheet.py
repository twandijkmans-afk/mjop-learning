#!/usr/bin/env python3
"""
export_review_sheet.py  (stap: NORMALIZED -> menselijke review)

Bouwt een Excel-werkblad (reports/review_batch1.xlsx) met alle
maintenance_actions en observations uit data/normalized/*.json die
requires_human_review=True hebben, zodat een mens ze in Excel kan
doorlopen i.p.v. in ruwe JSON.

Elke rij krijgt, naast de brontekst en de context (gekoppeld element,
bedragen), een REDEN kolom die letterlijk zegt waarom de pipeline deze post
niet vertrouwt (onbekende vocabulaire-term, kostenconflict, of een
afgeleide eenheidsprijs zonder letterlijke documentprijs - zie CLAUDE.md:
unit_cost_calculated is minder betrouwbaar dan een letterlijke prijs).

De mens vult drie lege kolommen in:
  BESLISSING (accept/edit/reject, zie CLAUDE.md), GECORRIGEERDE_WAARDE
  (alleen bij edit) en NOTITIES (vrije tekst, vooral nuttig bij reject).

scripts/apply_review.py leest het ingevulde bestand terug en schrijft
data/verified/*.json.

Gebruik:
    python3 scripts/export_review_sheet.py --batch batch_1
"""
import argparse
import glob
import json
import os

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

DECISION_OPTIONS = ["accept", "edit", "reject"]

ACTION_COLUMNS = [
    "document_id", "action_id", "element_id", "reden", "veld",
    "actie_original_value", "actie_huidige_normalized_value",
    "eenheid_original_value", "eenheid_huidige_normalized_value",
    "gekoppeld_element_type", "gekoppeld_element_locatie",
    "quantity_value", "planned_year",
    "total_cost_as_stated", "direct_cost_calculated", "unit_cost_calculated",
    "BESLISSING", "GECORRIGEERDE_WAARDE", "NOTITIES",
]

OBSERVATION_COLUMNS = [
    "document_id", "observation_id", "element_id", "reden",
    "defect_original_value", "defect_huidige_normalized_value",
    "description_value", "gekoppeld_element_type", "condition_score",
    "BESLISSING", "GECORRIGEERDE_WAARDE", "NOTITIES",
]


def action_review_reasons(action):
    reasons = []
    a = action.get("action") or {}
    u = action.get("unit") or {}
    if a.get("requires_human_review"):
        reasons.append("onbekende actieterm")
    if u.get("requires_human_review"):
        reasons.append("onbekende eenheid")
    if action.get("cost_conflict"):
        reasons.append("kostenconflict (vermeld totaal wijkt af van berekening)")
    uc_prov = ((action.get("unit_cost") or {}).get("provenance") or {})
    if uc_prov.get("source_confidence") == "low" and not reasons:
        reasons.append("afgeleide eenheidsprijs (geen letterlijke prijs in document)")
    return reasons or ["overig (zie provenance in data/normalized/)"]


def action_review_field(action):
    """Welk veld GECORRIGEERDE_WAARDE moet invullen. 'actie'/'eenheid' zijn
    de enige twee vocabulaire-paren op een maintenance_action; als geen van
    beide onbekend is (bijv. bij een kostenconflict of een afgeleide
    eenheidsprijs) is er geen vocabulaire-veld om te corrigeren - dan is
    alleen een BESLISSING (accept/reject) + NOTITIES zinvol, geen edit."""
    a = action.get("action") or {}
    u = action.get("unit") or {}
    if a.get("requires_human_review"):
        return "actie"
    if u.get("requires_human_review"):
        return "eenheid"
    return "geen"


def observation_review_reasons(obs):
    reasons = []
    d = obs.get("defect") or {}
    if d.get("requires_human_review"):
        reasons.append("onbekende gebreksterm")
    desc_prov = ((obs.get("description") or {}).get("provenance") or {})
    if (obs.get("description") or {}).get("requires_human_review") or desc_prov.get("source_confidence") == "low":
        reasons.append("lage extractie-confidence op omschrijving")
    return reasons or ["overig (zie provenance in data/normalized/)"]


def build_action_rows(rec, element_by_id):
    rows = []
    for action in rec.get("maintenance_actions", []):
        if not action.get("requires_human_review"):
            continue
        a = action.get("action") or {}
        u = action.get("unit") or {}
        el = element_by_id.get(action.get("element_id")) or {}
        el_type = (el.get("element_type") or {}).get("original_value")
        el_loc = (el.get("location") or {}).get("value")
        planned_year = action.get("planned_year")
        if isinstance(planned_year, dict):
            planned_year = planned_year.get("value")
        rows.append([
            rec.get("document_id"),
            action.get("action_id"),
            action.get("element_id"),
            "; ".join(action_review_reasons(action)),
            action_review_field(action),
            a.get("original_value"),
            a.get("normalized_value"),
            u.get("original_value"),
            u.get("normalized_value"),
            el_type,
            el_loc,
            (action.get("quantity") or {}).get("value"),
            planned_year,
            action.get("total_cost_as_stated"),
            action.get("direct_cost_calculated"),
            action.get("unit_cost_calculated"),
            None, None, None,
        ])
    return rows


def build_observation_rows(rec, element_by_id):
    rows = []
    for obs in rec.get("observations", []):
        if not obs.get("requires_human_review"):
            continue
        d = obs.get("defect") or {}
        el = element_by_id.get(obs.get("element_id")) or {}
        el_type = (el.get("element_type") or {}).get("original_value")
        rows.append([
            rec.get("document_id"),
            obs.get("observation_id"),
            obs.get("element_id"),
            "; ".join(observation_review_reasons(obs)),
            d.get("original_value"),
            d.get("normalized_value"),
            (obs.get("description") or {}).get("value"),
            el_type,
            (obs.get("condition_score") or {}).get("original_value"),
            None, None, None,
        ])
    return rows


def write_sheet(wb, title, columns, rows):
    ws = wb.create_sheet(title=title)
    ws.append(columns)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    for i, col in enumerate(columns, start=1):
        width = max(12, min(60, len(col) + 2))
        ws.column_dimensions[get_column_letter(i)].width = width
    if rows:
        decision_col = columns.index("BESLISSING") + 1
        dv = DataValidation(type="list", formula1='"accept,edit,reject"', allow_blank=True)
        ws.add_data_validation(dv)
        col_letter = get_column_letter(decision_col)
        dv.add(f"{col_letter}2:{col_letter}{len(rows) + 1}")
    ws.freeze_panes = "A2"
    return ws


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--batch", default="batch_1")
    ap.add_argument("--normalized-dir", default="data/normalized")
    ap.add_argument("--out", default="reports/review_batch1.xlsx")
    args = ap.parse_args()

    action_rows = []
    obs_rows = []
    for path in sorted(glob.glob(os.path.join(args.normalized_dir, "*.json"))):
        rec = json.load(open(path))
        element_by_id = {el.get("element_id"): el for el in rec.get("elements", [])}
        action_rows.extend(build_action_rows(rec, element_by_id))
        obs_rows.extend(build_observation_rows(rec, element_by_id))

    wb = Workbook()
    wb.remove(wb.active)
    write_sheet(wb, "maintenance_actions", ACTION_COLUMNS, action_rows)
    write_sheet(wb, "observations", OBSERVATION_COLUMNS, obs_rows)

    os.makedirs(os.path.dirname(args.out), exist_ok=True)
    wb.save(args.out)
    print(f"{len(action_rows)} maintenance_actions + {len(obs_rows)} observations -> {args.out}")


if __name__ == "__main__":
    main()
